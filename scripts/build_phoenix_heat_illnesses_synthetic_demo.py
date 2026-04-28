#!/usr/bin/env python3
"""
Builds `src/data/phoenixHeatIllnessesSyntheticDemo.json` from:
  External Datasets/Heat_Relief_Database_WithForecasting.xlsx :: Synthetic_Districts_DEMO

This keeps the app runtime simple (bundled JSON import) while letting the Excel
remain the canonical source of truth. Rows include both HISTORICAL and
FORECAST_2026 data; downstream UI can filter by Data_Type if desired.
"""

from __future__ import annotations

import datetime
import json
import re
from pathlib import Path

import openpyxl

XLSX_PATH = Path("External Datasets/Heat_Relief_Database_WithForecasting.xlsx")
SHEET = "Synthetic_Districts_DEMO"
OUT_PATH = Path("src/data/phoenixHeatIllnessesSyntheticDemo.json")


def iso_date(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()

    s = str(v).strip()
    m = re.match(r"^(\d{4}-\d{2}-\d{2})", s)
    if m:
        return m.group(1)

    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        mm, dd, yy = map(int, m.groups())
        return datetime.date(yy, mm, dd).isoformat()

    return s or None


def normalize_district(v) -> str:
    s = str(v or "").strip()
    if not s:
        return ""
    m = re.search(r"(?:^|\b)district\s*(\d+)\b", s, flags=re.I)
    if m:
        return m.group(1)
    m = re.match(r"^(\d+)\s*$", s)
    if m:
        return m.group(1)
    return s


def main() -> None:
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    ws = wb[SHEET]

    header_row_idx = None
    header = None
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if not row:
            continue
        row_norm = [str(c).strip() if c is not None else "" for c in row]
        if (
            "Week_Start" in row_norm
            and "Heat_Illness" in row_norm
            and "Count" in row_norm
            and "Council_District" in row_norm
        ):
            header_row_idx = idx
            header = row_norm
            break

    if header_row_idx is None or header is None:
        raise SystemExit(f"Could not find header row in {SHEET}")

    col_idx = {name: i for i, name in enumerate(header) if name}
    required = [
        "Week_Start",
        "Year",
        "Primary_Symptom",
        "Heat_Illness",
        "Council_District",
        "Count",
        "Data_Type",
        "Source",
    ]
    missing = [k for k in required if k not in col_idx]
    if missing:
        raise SystemExit(f"Missing columns in {SHEET}: {missing}. Found: {list(col_idx.keys())}")

    rows = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue

        def cell(k):
            return row[col_idx[k]]

        week = iso_date(cell("Week_Start"))
        illness = str(cell("Heat_Illness") or "").strip()
        district = normalize_district(cell("Council_District"))
        if not week or not illness or not district:
            continue

        try:
            year = int(cell("Year")) if cell("Year") is not None and str(cell("Year")).strip() != "" else None
        except Exception:
            year = None

        symptom = str(cell("Primary_Symptom") or "").strip()
        data_type = str(cell("Data_Type") or "").strip()
        source = str(cell("Source") or "").strip()

        raw_count = cell("Count")
        try:
            count = int(raw_count) if raw_count is not None and str(raw_count).strip() != "" else 0
        except Exception:
            try:
                count = int(float(raw_count))
            except Exception:
                count = 0

        rows.append(
            {
                "Week_Start": week,
                "Year": year,
                "Primary_Symptom": symptom,
                "Heat_Illness": illness,
                "Council_District": district,
                "Count": count,
                "Data_Type": data_type,
                "Source": source,
            }
        )

    rows.sort(
        key=lambda r: (
            r.get("Week_Start") or "",
            r.get("Year") or 0,
            r.get("Council_District") or "",
            r.get("Heat_Illness") or "",
            r.get("Primary_Symptom") or "",
            r.get("Data_Type") or "",
            r.get("Source") or "",
        )
    )

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(
        json.dumps(
            {
                "source": f"{XLSX_PATH.name}::{SHEET}",
                "note": "SYNTHETIC DATA — NOT REAL — FOR DEMO/PROTOTYPING ONLY. Includes HISTORICAL + FORECAST_2026 rows (see Data_Type).",
                "rows": rows,
            },
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )

    print(f"Wrote {OUT_PATH} ({len(rows)} rows)")


if __name__ == "__main__":
    main()

